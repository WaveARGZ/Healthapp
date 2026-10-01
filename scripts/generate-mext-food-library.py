"""日本食品標準成分表（八訂）増補2023年の第2章から検索用JSONを作る。

Usage: python3 scripts/generate-mext-food-library.py /path/to/第2章.xlsx
Requires: openpyxl (python3 -m pip install openpyxl)
"""

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


GROUPS = {
    "01": "ごはん・穀類",
    "02": "いも・でん粉",
    "03": "砂糖・甘味料",
    "04": "豆・大豆製品",
    "05": "種実類",
    "06": "野菜類",
    "07": "果物",
    "08": "きのこ類",
    "09": "海藻類",
    "10": "魚介類",
    "11": "肉類",
    "12": "卵類",
    "13": "乳製品",
    "14": "油脂類",
    "15": "菓子類",
    "16": "飲み物",
    "17": "調味料・香辛料",
    "18": "調理済み料理",
}


def nutrient(value):
    """推定値の括弧を外し、微量 (Tr) は表示精度に合わせ0とする。"""
    if value is None:
        return None
    text = str(value).strip().strip("()")
    if text == "Tr":
        return 0
    if text in ("", "-"):
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return round(number, 1) if number >= 0 else None


def main():
    if len(sys.argv) != 2:
        raise SystemExit("文部科学省の『第2章（データ）』xlsxを指定してください。")
    workbook = load_workbook(sys.argv[1], read_only=True, data_only=True)
    sheet = workbook["表全体"]
    foods = []
    skipped = []
    for row in sheet.iter_rows(min_row=13, values_only=True):
        group, code, _, raw_name = row[:4]
        if not code or not raw_name:
            continue
        macros = [nutrient(row[index]) for index in (6, 9, 12, 20)]
        if any(value is None for value in macros):
            skipped.append(code)
            continue
        name = re.sub(r"\s+", " ", raw_name).strip()
        foods.append({
            "id": f"mext-{code}",
            "name": name,
            "category": GROUPS[group],
            "aliases": "",
            "suggestedGrams": 100,
            "sourceDataset": "日本食品標準成分表（八訂）増補2023年",
            "foodCode": code,
            "sourceDescription": name,
            "per100g": dict(zip(("calories", "proteinG", "fatG", "carbsG"), macros)),
        })
    workbook.close()
    if len(foods) < 2500 or len({food["foodCode"] for food in foods}) != len(foods):
        raise ValueError("食品数または食品番号が想定外です。Excelの版を確認してください。")
    output = {
        "version": "bodymake-jp-2023-v1",
        "sourceName": "日本食品標準成分表（八訂）増補2023年",
        "sourceUrl": "https://www.mext.go.jp/a_menu/syokuhinseibun/mext_00001.html",
        "license": "文部科学省・出典明記による二次利用",
        "basis": "可食部100g当たり。Tr（微量）は0gとして表示。",
        "foods": foods,
    }
    target = Path("public/data/food-library-extended.json")
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(output, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    print(f"{len(foods)}件を生成。PFCが欠ける食品番号を除外: {', '.join(skipped)}")


if __name__ == "__main__":
    main()
